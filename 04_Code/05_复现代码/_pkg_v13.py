# -*- coding: utf-8 -*-
"""Round-13 data package upgrade:
 - new sheet 22_正文表格数值 : every cell of all 24 manuscript tables, long format
 - new sheet 23_图件作图数据 : Figure 1 residualized cell means (+counts) and
                              Figure 2 implied levels with recomputed SEs / 95% CIs
 - fix 20_平衡面板 note 1.197 -> 1.198 (recomputed value rounds to 1.198)
 - 00_索引 : new column H "本工作簿中的位置" + two new rows
"""
import os, io, sys, json, re, shutil
BASE = r'C:\Users\Administrator\WorkBuddy\2026-10-23-06-03-49'
os.chdir(BASE)
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P

FN = 'JAR档案_数据包_定稿_20261023.xlsx'
wb = openpyxl.load_workbook(FN)

HDR_FILL = PatternFill('solid', fgColor='EDF2F7')
SEC_FILL = PatternFill('solid', fgColor='FFF7E6')
SEC_FONT = Font(bold=True, size=12)
HDR_FONT = Font(bold=True)


def put(ws, r, c, v, bold=False, fill=False, wrap=False):
    cell = ws.cell(r, c, v)
    if bold:
        cell.font = HDR_FONT
    if fill:
        cell.fill = HDR_FILL
    if wrap:
        cell.alignment = Alignment(wrap_text=True, vertical='top')
    return cell


def put_hdr(ws, r, cols):
    for c, h in enumerate(cols, 1):
        put(ws, r, c, h, bold=True, fill=True)
    return r + 1


# ---------------------------------------------------------------- 22 ----------
doc = Document('JAR档案_Manuscript_定稿_20261023.docx')
body = doc.element.body
seq, tables = [], []     # (caption_context, table rows)
prev_nonempty = ''
table_no_hint = ''
for child in body.iterchildren():
    if isinstance(child, CT_P):
        p = Paragraph(child, doc)
        t = p.text.strip()
        if not t:
            continue
        prev_nonempty = t
        m = re.match(r'Table\s+(A\.\d+|\d+[A-Z]?)\.', t)
        if m:
            table_no_hint = 'Table ' + m.group(1)
    elif isinstance(child, CT_Tbl):
        tb = Table(child, doc)
        rows = [[c.text.strip().replace('\n', ' ') for c in r.cells] for r in tb.rows]
        seq.append((table_no_hint, prev_nonempty, rows))

ws = wb.create_sheet('22_正文表格数值')
r = 1
put(ws, r, 1, '22_正文表格数值 — 手稿 24 张表的逐格数值（长表；数据源 JAR档案_Manuscript_定稿_20261023.docx）', bold=True); r += 1
put(ws, r, 1, '每一行是一个表格单元：table = 所属表号，context = 紧邻表前的引导句（Panel 标题落在这一列），'
              'row_label = 该行第一格，col_index = 格在该行的序号（0 = 第一格），col_header = 仅当该行格数与表头行一致时按位置取到的表头，'
              'value_text = 原文，value_num = 可解析的数值（已去符号/千分位/百分号）。'
              '此表用于让数据编辑可检索稿件中任何一个数字，不需要打开 Word。', wrap=True)
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
ws.row_dimensions[r].height = 48
r += 2

cols = ['table', 'context', 'row_index', 'row_label', 'col_index', 'col_header', 'value_text', 'value_num']
put_hdr(ws, r, cols); r += 1


def to_num(s):
    if s is None:
        return None
    t = str(s).strip()
    txt = re.sub(r'[+−,%\s]', '', t)
    txt = txt.replace('−', '-')
    txt = re.sub(r'^[<>]=?', '', txt)
    txt = txt.split(' ')[0]
    # keep parenthesised-t statistics out of the numeric column
    m = re.fullmatch(r'-?\d+(?:\.\d+)?', txt)
    if m:
        try:
            return float(txt)
        except ValueError:
            return None
    return None


ncell = 0
for tno, ctx, rows in seq:
    if not rows:
        continue
    hdr_row = rows[0]
    for ri, rowvals in enumerate(rows):
        for ci, v in enumerate(rowvals):
            if v == '':
                continue
            ch = hdr_row[ci] if (len(rowvals) == len(hdr_row) and ri > 0) else ''
            put(ws, r, 1, tno)
            put(ws, r, 2, ctx[:200])
            put(ws, r, 3, ri)
            put(ws, r, 4, rowvals[0][:200])
            put(ws, r, 5, ci)
            put(ws, r, 6, ch[:120])
            put(ws, r, 7, v[:300])
            nv = to_num(v)
            if nv is not None:
                put(ws, r, 8, nv)
            r += 1
            ncell += 1
        r += 0
for i, w in enumerate([12, 52, 10, 40, 10, 30, 46, 12], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = 'A5'

# ---------------------------------------------------------------- 23 ----------
ws = wb.create_sheet('23_图件作图数据')
r = 1
put(ws, r, 1, '23_图件作图数据 — Figure 1 与 Figure 2 的作图坐标（DAN §8.17 存档值 + 本轮复算）', bold=True); r += 1
put(ws, r, 1, 'Figure 1 的 cell means / cell counts 来自 DAN §8.17 的存档块；cell-level SE 仍未存档，'
              '故 Figure 1 的 whisker 无法从本包重算（稿件图件由作者代码生成，方向与量级以 Table 8/10 为准）。'
              'Figure 2 的点位来自 §8.17 的交互设定，其 SE 已由本轮脚本 _fig2_check.py 从 18_主面板 复算补齐，'
              '95% 区间 = 点 ± 1.96 x SE（填补了 §8.17 "cell SE 未存档" 的缺口）。', wrap=True)
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
ws.row_dimensions[r].height = 62
r += 2

put(ws, r, 1, 'Figure 1 — 残差化 cell means（pair-firm+year FE 用于 pair 两行；firm+year FE 用于平衡面板行）', bold=True); r += 1
put_hdr(ws, r, ['Series', 'Clusters / firms', 'Cells', 'Clock year', 'Cell mean', 'Cell count', 'Cell SE (archived?)']); r += 1
f1 = [
    ('Joint clusters', 970, 2622, [1, 2, 3, 4, 5], [-0.202, -0.398, -0.269, 0.912, -0.612], [136, 223, 310, 915, 1038]),
    ('All other clusters', 6781, 16754, [1, 2, 3, 4, 5], [-0.078, -0.025, -0.046, 0.103, 0.078], [2390, 4461, 4525, 3547, 1831]),
    ('Balanced panel', 1203, None, [1, 2, 3, 4, 5], [-0.032, -0.038, -0.042, 0.019, 0.147], [2231, 3295, 3087, 2622, 1887]),
]
for name, cl, cells, yrs, means, cnts in f1:
    for y, mu, c in zip(yrs, means, cnts):
        put(ws, r, 1, name)
        put(ws, r, 2, cl)
        put(ws, r, 3, cells)
        put(ws, r, 4, y)
        put(ws, r, 5, mu)
        put(ws, r, 6, c)
        put(ws, r, 7, 'not archived')
        r += 1
r += 1

# Figure 2: archived levels from DAN 8.17 + recomputed SE from _fig2_check.py
num = json.load(open('_fig2_numbers.json', encoding='utf-8'))
nb = [v[0] for v in num['levels']['nb']]
nbse = [v[1] for v in num['levels']['nb']]
bb = [v[0] for v in num['levels']['b']]
bse = [v[1] for v in num['levels']['b']]
put(ws, r, 1, 'Figure 2 — 相对 year-one 的隐含水平（Y2-Y5 x balanced / non-balanced 交互设定；firm+year FE、控制、企业聚类）', bold=True); r += 1
put_hdr(ws, r, ['Series', 'Clock year', 'Implied level', 'SE (recomputed)', '95% CI low', '95% CI high', 'Source']); r += 1
for yr in range(1, 6):
    i = yr - 1
    for name, arr, se in (('Non-balanced firms', nb, nbse), ('Balanced firms', bb, bse)):
        put(ws, r, 1, name); put(ws, r, 2, yr); put(ws, r, 3, float(arr[i])); put(ws, r, 4, float(se[i]))
        put(ws, r, 5, float(arr[i]) - 1.96 * float(se[i])); put(ws, r, 6, float(arr[i]) + 1.96 * float(se[i]))
        put(ws, r, 7, 'year one base = 0' if yr == 1 else 'recomputed, _fig2_check.py')
        r += 1
r += 1
put(ws, r, 1, 'Figure 2 的落差（year four minus year five，正 = 披露下降）', bold=True); r += 1
put_hdr(ws, r, ['Series', 'Fall', 'SE', 't', 'MDE (80%)', '稿件写法', '备注']); r += 1
put(ws, r, 1, 'Non-balanced firms'); put(ws, r, 2, 1.198); put(ws, r, 3, 0.118); put(ws, r, 4, 10.19);
put(ws, r, 5, 0.33); put(ws, r, 6, '+1.198'); put(ws, r, 7, '复算 1.1979 → 三位小数 1.198（第五轮判定：图不动，文字以 1.198 为准）'); r += 1
put(ws, r, 1, 'Balanced firms'); put(ws, r, 2, -0.112); put(ws, r, 3, 0.127); put(ws, r, 4, -0.88);
put(ws, r, 5, 0.356); put(ws, r, 6, '-0.112'); put(ws, r, 7, '复算 -0.1120'); r += 1
r += 1
put(ws, r, 1, 'Pre-trend check（同年两组的隐含落差之差）', bold=True); r += 1
put_hdr(ws, r, ['Quantity', 'Estimate', 'SE', 't', 'MDE (80%)', '', '']); r += 1
for lbl, est_, se_, t_, mde in [
    ('Y2→Y3 difference-in-differences', -0.056, 0.151, -0.37, 0.423),
    ('Y3→Y4 difference-in-differences', -0.567, 0.162, -3.51, None),
    ('Y4→Y5 difference-in-differences (gap)', 1.310, 0.173, 7.56, None),
]:
    put(ws, r, 1, lbl); put(ws, r, 2, est_); put(ws, r, 3, se_); put(ws, r, 4, t_); put(ws, r, 5, mde); r += 1

for i, w in enumerate([30, 16, 12, 18, 16, 16, 30], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = 'A5'

# ------------------------------------------------- fix 20 & index -------------
ws20 = wb[[n for n in wb.sheetnames if n.startswith('20')][0]]
fixed = 0
for row in ws20.iter_rows():
    for cell in row:
        if isinstance(cell.value, str) and '1.197' in cell.value:
            cell.value = cell.value.replace('1.197', '1.198')
            fixed += 1

idx = wb['00_索引']
idx.cell(1, 8, '本工作簿中的位置')
POS = {
    '在本工作簿':
        {18: '18_主面板', 19: '19_Table8联合分解', 20: '20_平衡面板分解', 21: '21_企业层平衡检验',
         22: '22_正文表格数值', 23: '23_图件作图数据'},
}
for rr in range(2, idx.max_row + 1):
    first = idx.cell(rr, 1).value or ''
    if first.startswith('00_') or first.startswith('01_'):
        idx.cell(rr, 8, '本工作簿（见工作表名）')
    elif re.match(r'^\d\d_', first):
        no = int(first[:2])
        idx.cell(rr, 8, POS['在本工作簿'].get(no, '未随本工作簿发送 — 原始文件由作者持有，清单见 DAN §9.1 文件映射'))
    elif first.startswith('（未收录'):
        idx.cell(rr, 8, '未发送（太大 / 全文文本）')

idx.append(['22_正文表格数值', '档案自查（新增）', 'extracted from JAR档案_Manuscript_定稿_20261023.docx',
            ncell, 8, 'hand稿 24 张表的逐格数值长表：任一个正文数字都可在不必打开 Word 的前提下被检索到',
            '是', '22_正文表格数值'])
idx.append(['23_图件作图数据', '作图数据（新增）', 'DAN §8.17 + _fig2_check.py 复算',
            None, 7, 'Figure 1 的残差化 cell means 与 cell counts、Figure 2 的隐含水平及其复算 SE / 95% 区间、'
                     '落差与 pre-trend 检验；Figure 1 的 cell SE 仍未存档，已标注',
            '是', '23_图件作图数据'])
idx.column_dimensions['H'].width = 60

wb.save(FN)
print('saved', FN)
print('22 rows (data rows):', ncell)
print('20 fixed note cells:', fixed)
print('sheets:', wb.sheetnames)
print('index rows:', wb['00_索引'].max_row)
